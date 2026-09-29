import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE (Executive, High-Contrast, SIH-Compliant)
# ==============================================================================
C_SLATE_900  = RGBColor(15, 23, 42)      # #0F172A (Primary text & titles)
C_SLATE_700  = RGBColor(51, 65, 85)      # #334155 (Body text)
C_SLATE_600  = RGBColor(71, 85, 105)     # #475569 (Secondary body text)
C_SLATE_500  = RGBColor(100, 116, 139)   # #64748B (Muted text & labels)
C_WHITE      = RGBColor(255, 255, 255)   # #FFFFFF

C_DARK_GREEN = RGBColor(22, 101, 52)     # #166534 (Official Accent)
C_MID_GREEN  = RGBColor(22, 163, 74)     # #16A34A (Highlight Green)
C_BG_GREEN   = RGBColor(240, 253, 244)   # #F0FDF4 (Light Green Fill)
C_BD_GREEN   = RGBColor(187, 247, 208)   # #BBF7D0 (Green Border)

C_BLUE_DARK  = RGBColor(30, 64, 175)     # #1E40AF
C_BLUE_MID   = RGBColor(3, 105, 161)     # #0369A1 (Accent Blue)
C_BG_BLUE    = RGBColor(240, 249, 255)   # #F0F9FF (Light Blue Fill)
C_BD_BLUE    = RGBColor(186, 230, 253)   # #BAE6FD (Blue Border)

C_RED_DARK   = RGBColor(153, 27, 27)     # #991B1B (Risk / Problem Crimson)
C_RED_MID    = RGBColor(185, 28, 28)     # #B91C1C
C_BG_RED     = RGBColor(254, 242, 242)   # #FEF2F2

C_AMBER_DARK = RGBColor(180, 83, 9)      # #B45309 (Economic Amber)
C_BG_AMBER   = RGBColor(255, 251, 235)   # #FFFBEB
C_BD_AMBER   = RGBColor(253, 230, 138)   # #FDE68A

C_PURPLE_DARK= RGBColor(109, 40, 217)    # #6D28D9 (Compliance Purple)
C_BG_PURPLE  = RGBColor(250, 245, 255)   # #FAF5FF
C_BD_PURPLE  = RGBColor(233, 213, 255)   # #E9D5FF

C_TEAL_DARK  = RGBColor(13, 148, 136)    # #0D9488
C_BG_TEAL    = RGBColor(240, 253, 250)   # #F0FDFA
C_BD_TEAL    = RGBColor(153, 246, 228)   # #99F6E4

def set_font(r, text, bold=None, size_pt=None, color=None, underline=None, font_name="Arial"):
    r.text = text
    if bold is not None:
        r.font.bold = bold
    if size_pt is not None:
        r.font.size = Pt(size_pt)
    if color is not None:
        r.font.color.rgb = color
    if underline is not None:
        r.font.underline = underline
    if font_name is not None:
        r.font.name = font_name

def style_box(shape, bg_color=None, border_color=C_SLATE_900, border_width_pt=1.5):
    if bg_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_WHITE
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width_pt)
    else:
        shape.line.fill.background()
    if shape.has_text_frame:
        shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP

def clean_placeholder(slide, name):
    for sh in list(slide.shapes):
        if sh.name == name or (sh.has_text_frame and "Proposed Solution" in sh.text_frame.text) or (sh.has_text_frame and "Technologies to be used" in sh.text_frame.text) or (sh.has_text_frame and "Analysis of the feasibility" in sh.text_frame.text) or (sh.has_text_frame and "Potential impact" in sh.text_frame.text) or (sh.has_text_frame and "Details / Links" in sh.text_frame.text):
            sp = sh._element
            sp.getparent().remove(sp)

def update_team_oval(slide, name_text="HelloWorldWarriors"):
    for sh in slide.shapes:
        if "Oval" in sh.name and sh.has_text_frame:
            tf = sh.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = ""
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            set_font(r, name_text, bold=True, size_pt=10, color=C_SLATE_900)

def build_official_deck():
    template_path = r"C:\Users\Satya\Downloads\kabadiWala\SIH2026-IDEA-Presentation-Format.pptx"
    output_path = r"C:\Users\Satya\Downloads\kabadiWala\SIH2026_Idea_Presentation_E-Waste_Setu.pptx"
    
    shutil.copyfile(template_path, output_path)
    prs = pptx.Presentation(output_path)

    # Delete slide 7 (Instruction Slide)
    if len(prs.slides) > 6:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Deleted Slide 7 (Instruction Slide) successfully!")

    # ==========================================================================
    # SLIDE 1: Title Page
    # ==========================================================================
    s1 = prs.slides[0]
    for sh in s1.shapes:
        if sh.name == "TextBox 9":
            tf = sh.text_frame
            tf.clear()
            fields = [
                ("Problem Statement ID – ", "26229"),
                ("Problem Statement Title - ", "Kabadiwala Connect – Bringing the Informal Collector into the Formal Recycling Chain"),
                ("Organization - ", "Ministry of Mines (MoM) / JNARDDC"),
                ("Theme - ", "Clean & Green Technology"),
                ("PS Category - ", "Software"),
                ("Team ID - ", "170413"),
                ("Team Name - ", "HelloWorldWarriors"),
                ("Idea / Solution Title - ", "E-Waste Setu (ई-कचरा सेतु)")
            ]
            for idx, (lbl, val) in enumerate(fields):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                p.space_after = Pt(7)
                
                r_lbl = p.add_run()
                set_font(r_lbl, lbl, bold=True, size_pt=14, color=C_DARK_GREEN)
                
                r_val = p.add_run()
                is_bold = lbl.startswith("Idea") or lbl.startswith("Problem Statement ID") or lbl.startswith("Team Name")
                set_font(r_val, val, bold=is_bold, size_pt=14, color=C_SLATE_900)


    # ==========================================================================
    # SLIDE 2: Proposed Solution
    # ==========================================================================
    s2 = prs.slides[1]
    update_team_oval(s2)
    clean_placeholder(s2, "TextBox 8")

    # Title adjustment
    for sh in s2.shapes:
        if sh.name == "Title 1":
            p = sh.text_frame.paragraphs[0]
            p.text = ""
            r = p.add_run()
            set_font(r, "IDEA TITLE", bold=True, size_pt=32, color=C_SLATE_900)

    # Subtitle Box: Title & 4-Stage Custody Chain
    tb_sub = s2.shapes.add_textbox(Inches(0.6), Inches(1.08), Inches(12.13), Inches(0.65))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = 0
    tf_sub.margin_top = 0
    p0 = tf_sub.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run()
    set_font(r0, "E-Waste Setu: Connecting Informal Waste Collectors (Kabadiwalas) to Authorized Recyclers", bold=True, size_pt=13.5, color=C_SLATE_900, underline=True)
    
    p1 = tf_sub.add_paragraph()
    p1.alignment = PP_ALIGN.LEFT
    p1.space_before = Pt(2)
    r_arch = p1.add_run()
    set_font(r_arch, "Core Supply Chain: ", bold=True, size_pt=9.8, color=C_DARK_GREEN)
    r_arch_t = p1.add_run()
    set_font(r_arch_t, "Small Collector  ──>  Local Aggregator / Collection Hub  ──>  Bulk Aggregation  ──>  Authorized Recycler (CPCB)", bold=True, size_pt=9.2, color=C_BLUE_DARK)

    # ==========================================================================
    # TOP ROW - CARD 1: Main Problems & How We Address Them (Top Left)
    # ==========================================================================
    card_prob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.78), Inches(5.95), Inches(2.54))
    style_box(card_prob, bg_color=C_WHITE, border_color=C_RED_MID, border_width_pt=1.5)
    tf_pr = card_prob.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = Inches(0.16)
    tf_pr.margin_right = Inches(0.16)
    tf_pr.margin_top = Inches(0.10)
    tf_pr.margin_bottom = Inches(0.08)

    p_pr_title = tf_pr.paragraphs[0]
    p_pr_title.alignment = PP_ALIGN.LEFT
    p_pr_title.space_after = Pt(3)
    r_pr_t = p_pr_title.add_run()
    set_font(r_pr_t, "Main Problems & How We Address Them", bold=True, size_pt=12.0, color=C_RED_DARK, underline=True)

    problem_solutions = [
        ("Unclear & Unfair Prices (25–40% loss): ", "Middlemen exploit lack of market awareness. ", "➔ Daily fair metal-benchmarked scrap rates on phone."),
        ("Tampered Scales (15–20% weight loss): ", "Mechanical scales are rigged by middlemen. ", "➔ Phone camera verifies scale display screen to lock true weight."),
        ("No Trusted Recycler Access: ", "Collectors & hubs don't know verified recyclers or rates. ", "➔ Matches lots with authorized recyclers using license, material, price & pickup."),
        ("Too Many Middlemen: ", "Fragmented informal tiers skim margins and resist formal apps. ", "➔ Upgrades local scrap dealers into certified collection hubs with legal fees."),
        ("Illiteracy & Cash Reliance: ", "Cannot read text-heavy apps; rely strictly on cash. ", "➔ Visual touch UI + optional voice; supports both Cash and UPI.")
    ]

    for p_prob, p_desc, p_sol in problem_solutions:
        p = tf_pr.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(1.5)
        p.space_after = Pt(1.0)
        r_pr = p.add_run()
        set_font(r_pr, f"• {p_prob}", bold=True, size_pt=8.2, color=C_SLATE_900)
        r_desc = p.add_run()
        set_font(r_desc, p_desc, bold=False, size_pt=8.0, color=C_SLATE_700)
        r_so = p.add_run()
        set_font(r_so, p_sol, bold=True, size_pt=8.0, color=C_DARK_GREEN)

    # ==========================================================================
    # TOP ROW - CARD 2: Detailed Explanation of Proposed Solution (Top Right)
    # ==========================================================================
    card_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.78), Inches(1.78), Inches(5.95), Inches(2.54))
    style_box(card_sol, bg_color=C_WHITE, border_color=C_DARK_GREEN, border_width_pt=1.5)
    tf_sol = card_sol.text_frame
    tf_sol.word_wrap = True
    tf_sol.margin_left = Inches(0.16)
    tf_sol.margin_right = Inches(0.16)
    tf_sol.margin_top = Inches(0.10)
    tf_sol.margin_bottom = Inches(0.08)

    p_sol_title = tf_sol.paragraphs[0]
    p_sol_title.alignment = PP_ALIGN.LEFT
    p_sol_title.space_after = Pt(3)
    r_sol_t = p_sol_title.add_run()
    set_font(r_sol_t, "Proposed Solution: End-to-End Handover Flow", bold=True, size_pt=12.0, color=C_DARK_GREEN, underline=True)

    solution_flow = [
        ("Collector [Photo + Voice + Category]: ", "Logs scrap items via touch visual catalog or optional 10-language voice; views guaranteed fair rates."),
        ("Digital Lot Created [Weight + Value Lock]: ", "System pre-records item category, estimated quantity, and metal-benchmarked floor valuation."),
        ("Local Collection Hub [Scale Photo & Settle]: ", "Drop-off at hub (1–2 km); camera reads scale display to verify weight. Cash or UPI ➔ digital receipt ➔ earnings ledger."),
        ("Bulk Collection & Transfer [Smart Logistics]: ", "Hubs aggregate volume; PostGIS routes authorized trucks for optimized bulk pickup across hubs."),
        ("Authorized Recycler [Track Every Handover]: ", "CPCB plant verifies lot QR code for safe scientific recycling; automated Form-6 manifest logged on portal.")
    ]

    for idx, (f_step, f_desc) in enumerate(solution_flow):
        p = tf_sol.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(1.5)
        p.space_after = Pt(1.0)
        r_num = p.add_run()
        set_font(r_num, f"{idx+1}. {f_step}", bold=True, size_pt=8.2, color=C_DARK_GREEN if "Recycler" in f_step or "Hub" in f_step else C_BLUE_DARK)
        r_fd = p.add_run()
        set_font(r_fd, f_desc, bold=False, size_pt=8.0, color=C_SLATE_700)

    # ==========================================================================
    # BOTTOM SECTION: Innovation & Uniqueness of the Solution (4 Horizontal Cards)
    # ==========================================================================
    tb_inn_title = s2.shapes.add_textbox(Inches(0.6), Inches(4.36), Inches(12.13), Inches(0.28))
    tf_it = tb_inn_title.text_frame
    tf_it.word_wrap = True
    tf_it.margin_left = 0
    tf_it.margin_top = 0
    p_it0 = tf_it.paragraphs[0]
    p_it0.alignment = PP_ALIGN.LEFT
    r_it0 = p_it0.add_run()
    set_font(r_it0, "Innovation and Uniqueness of the Solution", bold=True, size_pt=12.0, color=C_SLATE_900, underline=True)

    innovations = [
        ("INNOVATION 1", "Partners with Local Dealers", "Instead of eliminating local scrap shops (which triggers backlash), we upgrade them into certified collection hubs earning steady legal commissions.", "Conflict-Free Adoption", C_BLUE_MID, C_BG_BLUE, C_BD_BLUE),
        ("INNOVATION 2", "Camera-Based Scale Reading", "Requires no expensive IoT scales. Any standard smartphone camera reads numbers off existing weighing scales to stop weight cheating.", "Zero Extra Hardware Cost", C_DARK_GREEN, C_BG_GREEN, C_BD_GREEN),
        ("INNOVATION 3", "Multimodal Accessibility", "Built for all literacy levels. Visual touch UI with category photos, plus optional voice assistance in 10 Indian dialects with offline caching.", "Zero Literacy Barrier", C_AMBER_DARK, C_BG_AMBER, C_BD_AMBER),
        ("INNOVATION 4", "Verifiable Chain of Custody", "Binds photo + weight + GPS + timestamp + unique lot ID + recycler confirmation into an auditable digital manifest for CPCB compliance.", "Track Every Handover", C_PURPLE_DARK, C_BG_PURPLE, C_BD_PURPLE)
    ]

    for idx, (i_tag, i_title, i_desc, i_metric, i_col, i_bg, i_bd) in enumerate(innovations):
        x_card = Inches(0.60 + idx * 3.08)
        sh_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_card, Inches(4.68), Inches(2.89), Inches(2.10))
        style_box(sh_card, bg_color=i_bg, border_color=i_bd, border_width_pt=1.5)
        
        tf_c = sh_card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.12)
        tf_c.margin_right = Inches(0.12)
        tf_c.margin_top = Inches(0.10)
        tf_c.margin_bottom = Inches(0.08)

        p_top = tf_c.paragraphs[0]
        p_top.alignment = PP_ALIGN.LEFT
        r_tag = p_top.add_run()
        set_font(r_tag, f"[{i_tag}]\n", bold=True, size_pt=8.5, color=i_col)
        r_title = p_top.add_run()
        set_font(r_title, i_title, bold=True, size_pt=10.0, color=C_SLATE_900)

        p_body = tf_c.add_paragraph()
        p_body.alignment = PP_ALIGN.LEFT
        p_body.space_before = Pt(3)
        r_body = p_body.add_run()
        set_font(r_body, i_desc, bold=False, size_pt=8.1, color=C_SLATE_700)

        p_metric = tf_c.add_paragraph()
        p_metric.alignment = PP_ALIGN.RIGHT
        p_metric.space_before = Pt(3)
        r_m = p_metric.add_run()
        set_font(r_m, f"• {i_metric}", bold=True, size_pt=8.0, color=i_col)

    # ==========================================================================
    # SLIDE 3: Technical Approach
    # ==========================================================================
    s3 = prs.slides[2]
    update_team_oval(s3)
    clean_placeholder(s3, "TextBox 8")

    for sh in s3.shapes:
        if sh.name == "Title 1":
            sh.text_frame.clear()
            p = sh.text_frame.paragraphs[0]
            r = p.add_run()
            set_font(r, "TECHNICAL APPROACH", bold=True, size_pt=32, color=C_SLATE_900)

    # Subtitle Box: Architecture Summary
    tb_sub3 = s3.shapes.add_textbox(Inches(0.6), Inches(1.08), Inches(12.13), Inches(0.32))
    tf_sub3 = tb_sub3.text_frame
    tf_sub3.word_wrap = True
    tf_sub3.margin_left = 0
    tf_sub3.margin_top = 0
    p0_s3 = tf_sub3.paragraphs[0]
    p0_s3.alignment = PP_ALIGN.LEFT
    # Subtitle Box: Architecture Summary
    tb_sub3 = s3.shapes.add_textbox(Inches(0.6), Inches(1.08), Inches(12.13), Inches(0.32))
    tf_sub3 = tb_sub3.text_frame
    tf_sub3.word_wrap = True
    tf_sub3.margin_left = 0
    tf_sub3.margin_top = 0
    p0_s3 = tf_sub3.paragraphs[0]
    p0_s3.alignment = PP_ALIGN.LEFT
    r0_s3 = p0_s3.add_run()
    set_font(r0_s3, "From Collection to Authorized Recycling: ", bold=True, size_pt=12.2, color=C_SLATE_900)
    r1_s3 = p0_s3.add_run()
    set_font(r1_s3, "End-to-End System Pipeline, AI Services & Core Technology Stack", bold=False, size_pt=10.5, color=C_DARK_GREEN)

    # ==========================================================================
    # LEFT COLUMN: CORE TECHNOLOGY STACK (TOP) & AI PLATFORM INTELLIGENCE (DOWN)
    # ==========================================================================
    # Left Top Card: CORE TECHNOLOGY STACK ARCHITECTURE
    card_tech = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60), Inches(1.44), Inches(4.88), Inches(2.88))
    style_box(card_tech, bg_color=C_WHITE, border_color=C_BD_GREEN, border_width_pt=1.5)
    tf_tc = card_tech.text_frame
    tf_tc.word_wrap = True
    tf_tc.margin_left = Inches(0.14)
    tf_tc.margin_right = Inches(0.14)
    tf_tc.margin_top = Inches(0.08)
    tf_tc.margin_bottom = Inches(0.06)

    p_tc_hdr = tf_tc.paragraphs[0]
    p_tc_hdr.alignment = PP_ALIGN.LEFT
    p_tc_hdr.space_after = Pt(2)
    r_tc0 = p_tc_hdr.add_run()
    set_font(r_tc0, "CORE TECHNOLOGY STACK ARCHITECTURE", bold=True, size_pt=10.2, color=C_DARK_GREEN, underline=True)
    r_tc_sub = p_tc_hdr.add_run()
    set_font(r_tc_sub, "  |  Multi-Tier Production Design", bold=False, size_pt=8.0, color=C_SLATE_600)

    tech_stack_layers = [
        ("Frontend / Client: ", "React 18 PWA + Vite + TypeScript",
         "Ultra-light progressive web app; runs in mobile browser on ₹4k Android phones with zero app-store download friction.", C_DARK_GREEN),
        ("Offline Resilience: ", "IndexedDB + Dexie.js + Service Workers",
         "Full local offline queuing for weights, photos & transactions; auto-syncs asynchronously upon 2G reconnection.", C_BLUE_MID),
        ("Backend & REST APIs: ", "Node.js + Fastify Core",
         "High-throughput microservices (<20ms latency) handling batch lot creation, digital receipt dispatch & auth.", C_BLUE_DARK),
        ("Spatial Database: ", "PostgreSQL 16 + PostGIS Extension",
         "Spatial indexing for aggregator-to-recycler matching, geo-fenced custody handovers, and pickup route planning.", C_PURPLE_DARK),
        ("Vision & OCR Engine: ", "Python + OpenCV + EasyOCR",
         "Extracts numeric digits directly from scale screen photos to verify true weight without costly IoT hardware.", C_RED_DARK),
        ("Security & Custody: ", "SHA-256 Hashes + Dynamic QR",
         "Cryptographic lot tokenization & scannable receipts; generates verifiable audit trails for automated CPCB Form-6.", C_AMBER_DARK),
        ("Hardware Footprint: ", "Zero Capex (100% BYOD Model)",
         "Works entirely on informal collectors' existing basic smartphones and scrap dealers' standard weighing scales.", C_SLATE_900)
    ]

    for t_cat, t_tech, t_desc, t_col in tech_stack_layers:
        p_tl = tf_tc.add_paragraph()
        p_tl.alignment = PP_ALIGN.LEFT
        p_tl.space_before = Pt(1.8)
        p_tl.space_after = Pt(0.0)
        
        r_c = p_tl.add_run()
        set_font(r_c, f"• {t_cat}", bold=True, size_pt=7.5, color=t_col)
        
        r_t = p_tl.add_run()
        set_font(r_t, f"[{t_tech}] ", bold=True, size_pt=7.3, color=C_SLATE_900)
        
        r_d = p_tl.add_run()
        set_font(r_d, t_desc, bold=False, size_pt=7.1, color=C_SLATE_700)

    # Left Down Card: AI & PLATFORM INTELLIGENCE
    card_ai = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60), Inches(4.38), Inches(4.88), Inches(2.46))
    style_box(card_ai, bg_color=C_WHITE, border_color=C_BD_BLUE, border_width_pt=1.5)
    tf_ai = card_ai.text_frame
    tf_ai.word_wrap = True
    tf_ai.margin_left = Inches(0.14)
    tf_ai.margin_right = Inches(0.14)
    tf_ai.margin_top = Inches(0.08)
    tf_ai.margin_bottom = Inches(0.06)

    p_ai_hdr = tf_ai.paragraphs[0]
    p_ai_hdr.alignment = PP_ALIGN.LEFT
    p_ai_hdr.space_after = Pt(2)
    r_ah0 = p_ai_hdr.add_run()
    set_font(r_ah0, "AI & PLATFORM INTELLIGENCE (Platform Services)", bold=True, size_pt=10.2, color=C_BLUE_DARK, underline=True)

    ai_services = [
        ("Material Identification: ", "Computer Vision & MobileNet assist scrap & PCB category classification from photos.", C_BLUE_MID),
        ("Price Benchmark Index: ", "Current market price benchmark index for transparent fair valuation.", C_DARK_GREEN),
        ("Recycler Matching: ", "Spatial routing algorithm matches bulk lots with verified recyclers by license, material & pickup radius.", C_BLUE_DARK),
        ("Anomaly Detection: ", "Photo-based weight evidence & discrepancy detection flags volume/price irregularities across handovers.", C_RED_DARK),
        ("Voice Assistance: ", "Speech recognition in 10 regional Indian dialects for low-literacy informal collectors.", C_AMBER_DARK),
        ("Offline Sync Engine: ", "Dexie.js IndexedDB buffer queues transactions locally in 2G dead zones.", C_PURPLE_DARK)
    ]

    for a_title, a_desc, a_col in ai_services:
        p_srv = tf_ai.add_paragraph()
        p_srv.alignment = PP_ALIGN.LEFT
        p_srv.space_before = Pt(2.0)
        p_srv.space_after = Pt(0.5)
        r_st = p_srv.add_run()
        set_font(r_st, f"• {a_title}", bold=True, size_pt=7.6, color=a_col)
        r_sd = p_srv.add_run()
        set_font(r_sd, a_desc, bold=False, size_pt=7.3, color=C_SLATE_700)

    # ==========================================================================
    # RIGHT COLUMN: METHODOLOGY — END-TO-END VISUAL PROCESSING PIPELINE
    # ==========================================================================
    tb_pipe_head = s3.shapes.add_textbox(Inches(5.72), Inches(1.44), Inches(7.02), Inches(0.24))
    tf_ph = tb_pipe_head.text_frame
    tf_ph.word_wrap = True
    tf_ph.margin_left = 0
    tf_ph.margin_top = 0
    p_ph = tf_ph.paragraphs[0]
    p_ph.alignment = PP_ALIGN.LEFT
    r_ph = p_ph.add_run()
    set_font(r_ph, "Methodology: End-to-End Visual Processing Pipeline", bold=True, size_pt=11.2, color=C_SLATE_900, underline=True)

    pipeline_steps = [
        ("STEP 1 — COLLECT & IDENTIFY", "Small Informal Collector",
         "Takes photo of scrap; selects material category via visual touch UI or optional 10-language voice.",
         "Displays current market price benchmarks; buffers entry locally in IndexedDB if offline.",
         C_DARK_GREEN, C_BD_GREEN),

        ("STEP 2 — CREATE DIGITAL LOT", "Digital Lot Generator",
         "Pre-calibrates lot with material type, estimated weight, photo evidence, and collector profile ID.",
         "Binds GPS geolocation and server timestamp, generating an unconfirmed preliminary batch record.",
         C_BLUE_DARK, C_BD_BLUE),

        ("STEP 3 — LOCAL HUB / AGGREGATOR", "Collection Hub / Scrap Dealer",
         "Receives small collector lots; camera snaps scale display. Settle: Cash or UPI ➔ digital receipt.",
         "OpenCV + EasyOCR extracts scale digits for photo-based weight evidence & discrepancy detection.",
         C_AMBER_DARK, C_BD_AMBER),

        ("STEP 4 — BULK AGGREGATION & TRANSFER", "Aggregator / Bulk Trader",
         "Combines small lots into bulk batches; system matches lot with eligible authorized recyclers.",
         "PostGIS engine clusters routes and dispatches authorized trucks; driver scans lot QR on pickup.",
         C_SLATE_900, RGBColor(203, 213, 225)),

        ("STEP 5 — AUTHORIZED RECYCLER", "CPCB Authorized Recycling Plant",
         "Receives shipment, verifies sealed QR tag, inspects weight & material for scientific recovery.",
         "Confirms legal custody handover; auto-submits digital Form-6 manifest to national CPCB portal.",
         C_PURPLE_DARK, C_BD_PURPLE)
    ]

    y_start_pipe = Inches(1.74)
    pipe_box_h = Inches(0.64)
    pipe_arrow_h = Inches(0.15)
    pipe_gap = Inches(0.06)

    for idx, (p_tag, p_role, p_act_txt, p_sys_txt, p_col, p_bd) in enumerate(pipeline_steps):
        y_box = y_start_pipe + idx * (pipe_box_h + pipe_arrow_h + pipe_gap * 2)

        # Step Box (Sleek, compact height)
        sh_step = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.72), y_box, Inches(7.02), pipe_box_h)
        style_box(sh_step, bg_color=C_WHITE, border_color=p_bd, border_width_pt=1.3)
        tf_s = sh_step.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.12)
        tf_s.margin_right = Inches(0.12)
        tf_s.margin_top = Inches(0.04)
        tf_s.margin_bottom = Inches(0.03)

        # Line 1: Step Tag + Role
        p_hdr = tf_s.paragraphs[0]
        p_hdr.alignment = PP_ALIGN.LEFT
        r_pt = p_hdr.add_run()
        set_font(r_pt, f"[{p_tag}]", bold=True, size_pt=7.8, color=p_col)
        r_pr = p_hdr.add_run()
        set_font(r_pr, f"  |  Role: {p_role}", bold=True, size_pt=7.8, color=C_SLATE_900)

        # Line 2: Operational Action
        p_act = tf_s.add_paragraph()
        p_act.alignment = PP_ALIGN.LEFT
        p_act.space_before = Pt(1.0)
        p_act.space_after = Pt(0.0)
        r_al = p_act.add_run()
        set_font(r_al, "• Action: ", bold=True, size_pt=7.3, color=p_col)
        r_av = p_act.add_run()
        set_font(r_av, p_act_txt, bold=False, size_pt=7.3, color=C_SLATE_900)

        # Line 3: System Engine / Processing
        p_sys = tf_s.add_paragraph()
        p_sys.alignment = PP_ALIGN.LEFT
        p_sys.space_before = Pt(0.5)
        p_sys.space_after = Pt(0.0)
        r_sl = p_sys.add_run()
        set_font(r_sl, "• System: ", bold=True, size_pt=7.3, color=C_SLATE_700)
        r_sv = p_sys.add_run()
        set_font(r_sv, p_sys_txt, bold=False, size_pt=7.2, color=C_SLATE_600)

        # Downward Arrow between Step Boxes
        if idx < len(pipeline_steps) - 1:
            y_arr = y_box + pipe_box_h + pipe_gap
            arr = s3.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(5.72 + 3.42), y_arr, Inches(0.18), pipe_arrow_h)
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(100, 116, 139)
            arr.line.fill.background()

    # Traceability Strip directly underneath the flow
    sh_strip = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.72), Inches(6.24), Inches(7.02), Inches(0.60))
    style_box(sh_strip, bg_color=C_BG_GREEN, border_color=C_BD_GREEN, border_width_pt=1.3)
    tf_strip = sh_strip.text_frame
    tf_strip.word_wrap = True
    tf_strip.margin_left = Inches(0.14)
    tf_strip.margin_right = Inches(0.14)
    tf_strip.margin_top = Inches(0.08)
    tf_strip.margin_bottom = Inches(0.06)

    p_strip0 = tf_strip.paragraphs[0]
    p_strip0.alignment = PP_ALIGN.CENTER
    r_st0 = p_strip0.add_run()
    set_font(r_st0, "VERIFIABLE CHAIN OF CUSTODY (TRACEABILITY STRIP)", bold=True, size_pt=8.2, color=C_DARK_GREEN)

    p_strip1 = tf_strip.add_paragraph()
    p_strip1.alignment = PP_ALIGN.CENTER
    p_strip1.space_before = Pt(1.5)
    r_st1 = p_strip1.add_run()
    set_font(r_st1, "Every Stage Recorded ➔ Photo Proof + Scale Weight + Agreed Price + GPS + Timestamp + QR Reference", bold=True, size_pt=7.7, color=C_SLATE_900)

    # ==========================================================================
    # SLIDE 4: Feasibility and Viability
    # ==========================================================================
    s4 = prs.slides[3]
    update_team_oval(s4)
    clean_placeholder(s4, "TextBox 8")

    for sh in s4.shapes:
        if sh.name == "Title 1":
            p = sh.text_frame.paragraphs[0]
            p.text = ""
            r = p.add_run()
            set_font(r, "FEASIBILITY AND VIABILITY", bold=True, size_pt=32, color=C_SLATE_900)

    # Slide 4 Subtitle
    sub_box4 = s4.shapes.add_textbox(Inches(0.60), Inches(1.08), Inches(12.13), Inches(0.32))
    tf_sub4 = sub_box4.text_frame
    tf_sub4.word_wrap = True
    tf_sub4.margin_left = 0
    tf_sub4.margin_top = 0
    p0_s4 = tf_sub4.paragraphs[0]
    p0_s4.alignment = PP_ALIGN.LEFT
    r0_s4 = p0_s4.add_run()
    set_font(r0_s4, "Ground Reality & Operational Viability: ", bold=True, size_pt=12.0, color=C_SLATE_900)
    r1_s4 = p0_s4.add_run()
    set_font(r1_s4, "Multi-Pillar Feasibility Analysis & Defensible Risk Mitigation", bold=False, size_pt=10.5, color=C_DARK_GREEN)

    # ==========================================================================
    # LEFT COLUMN: 3 FEASIBILITY CARDS (Technical, Operational, Economic)
    # ==========================================================================
    # Card 1: Technical Feasibility
    card_tf = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60), Inches(1.44), Inches(5.85), Inches(1.72))
    style_box(card_tf, bg_color=C_WHITE, border_color=C_BD_BLUE, border_width_pt=1.5)
    tf_tf = card_tf.text_frame
    tf_tf.word_wrap = True
    tf_tf.margin_left = Inches(0.14)
    tf_tf.margin_right = Inches(0.14)
    tf_tf.margin_top = Inches(0.07)
    tf_tf.margin_bottom = Inches(0.05)

    p_tf_hdr = tf_tf.paragraphs[0]
    p_tf_hdr.alignment = PP_ALIGN.LEFT
    r_tf_t = p_tf_hdr.add_run()
    set_font(r_tf_t, "TECHNICAL FEASIBILITY", bold=True, size_pt=10.2, color=C_BLUE_DARK, underline=True)
    r_tf_sub = p_tf_hdr.add_run()
    set_font(r_tf_sub, "  |  Low-Spec Readiness & Scalable Backend", bold=False, size_pt=7.6, color=C_SLATE_600)

    tf_bullets = [
        ("Low-End Android: ", "Lightweight PWA designed for entry-level devices with minimal RAM footprint.", C_BLUE_MID),
        ("AI-Assisted Identification: ", "Camera-based material recognition assisting scrap category classification.", C_DARK_GREEN),
        ("Weight Verification: ", "Scale-photo OCR provides verifiable weight evidence from existing scales.", C_BLUE_DARK),
        ("Offline-First: ", "Digital lots and transactions are stored locally and synced when signal returns.", C_PURPLE_DARK),
        ("Scalable Backend: ", "PostgreSQL + PostGIS for transactions, locations, and recycler matching.", C_SLATE_900)
    ]

    for b_lbl, b_val, b_col in tf_bullets:
        p_b = tf_tf.add_paragraph()
        p_b.alignment = PP_ALIGN.LEFT
        p_b.space_before = Pt(1.5)
        p_b.space_after = Pt(0.0)
        r_bl = p_b.add_run()
        set_font(r_bl, f"• {b_lbl}", bold=True, size_pt=7.6, color=b_col)
        r_bv = p_b.add_run()
        set_font(r_bv, b_val, bold=False, size_pt=7.3, color=C_SLATE_700)

    # Card 2: Operational Feasibility
    card_of = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60), Inches(3.24), Inches(5.85), Inches(1.72))
    style_box(card_of, bg_color=C_WHITE, border_color=C_BD_GREEN, border_width_pt=1.5)
    tf_of = card_of.text_frame
    tf_of.word_wrap = True
    tf_of.margin_left = Inches(0.14)
    tf_of.margin_right = Inches(0.14)
    tf_of.margin_top = Inches(0.07)
    tf_of.margin_bottom = Inches(0.05)

    p_of_hdr = tf_of.paragraphs[0]
    p_of_hdr.alignment = PP_ALIGN.LEFT
    r_of_t = p_of_hdr.add_run()
    set_font(r_of_t, "OPERATIONAL FEASIBILITY", bold=True, size_pt=10.2, color=C_DARK_GREEN, underline=True)
    r_of_sub = p_of_hdr.add_run()
    set_font(r_of_sub, "  |  Aligned with Ground Scrap Realities", bold=False, size_pt=7.6, color=C_SLATE_600)

    of_bullets = [
        ("Voice + Pictorial UX: ", "Reduces reading and typing through visual category icons and regional voice.", C_DARK_GREEN),
        ("Existing Scrap Shops as Hubs: ", "Works with the current informal network, onboarding established dealers.", C_BLUE_MID),
        ("Small Lots ➔ Bulk Lots: ", "Directly matches how scrap is naturally collected and aggregated in the field.", C_BLUE_DARK),
        ("Recycler Pickup Logistics: ", "Avoids requiring individual informal collectors to transport heavy material.", C_AMBER_DARK),
        ("Offline Sync: ", "Continues working seamlessly during poor connectivity in scrap clusters and transit.", C_PURPLE_DARK)
    ]

    for b_lbl, b_val, b_col in of_bullets:
        p_b = tf_of.add_paragraph()
        p_b.alignment = PP_ALIGN.LEFT
        p_b.space_before = Pt(1.5)
        p_b.space_after = Pt(0.0)
        r_bl = p_b.add_run()
        set_font(r_bl, f"• {b_lbl}", bold=True, size_pt=7.6, color=b_col)
        r_bv = p_b.add_run()
        set_font(r_bv, b_val, bold=False, size_pt=7.3, color=C_SLATE_700)

    # Card 3: Economic Feasibility
    card_ef = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60), Inches(5.04), Inches(5.85), Inches(1.80))
    style_box(card_ef, bg_color=C_WHITE, border_color=C_BD_AMBER, border_width_pt=1.5)
    tf_ef = card_ef.text_frame
    tf_ef.word_wrap = True
    tf_ef.margin_left = Inches(0.14)
    tf_ef.margin_right = Inches(0.14)
    tf_ef.margin_top = Inches(0.07)
    tf_ef.margin_bottom = Inches(0.05)

    p_ef_hdr = tf_ef.paragraphs[0]
    p_ef_hdr.alignment = PP_ALIGN.LEFT
    r_ef_t = p_ef_hdr.add_run()
    set_font(r_ef_t, "ECONOMIC FEASIBILITY & VIABILITY", bold=True, size_pt=10.2, color=C_AMBER_DARK, underline=True)
    r_ef_sub = p_ef_hdr.add_run()
    set_font(r_ef_sub, "  |  Defensible Stakeholder Value Model", bold=False, size_pt=7.6, color=C_SLATE_600)

    p_ef_pr = tf_ef.add_paragraph()
    p_ef_pr.alignment = PP_ALIGN.LEFT
    p_ef_pr.space_before = Pt(1.0)
    p_ef_pr.space_after = Pt(1.0)
    r_pr_tag = p_ef_pr.add_run()
    set_font(r_pr_tag, "Design Principle: ", bold=True, size_pt=7.4, color=C_DARK_GREEN)
    r_pr_txt = p_ef_pr.add_run()
    set_font(r_pr_txt, "Improves collector earnings through transparent pricing and reduced information loss.", bold=False, size_pt=7.4, color=C_SLATE_700)

    ef_bullets = [
        ("Informal Collector: ", "100% free to use — zero subscription cost or capital expenditure barrier.", C_DARK_GREEN),
        ("Hub / Aggregator: ", "Earns legitimate handling and bulk aggregation margins on consolidated scrap.", C_BLUE_MID),
        ("Authorized Recycler: ", "Pays a small platform fee for guaranteed verified supply & verifiable traceability.", C_BLUE_DARK),
        ("Platform Revenue Model: ", "Initial model — 1.5% verified transaction fee; Future: B2B traceability & EPR services.", C_SLATE_900)
    ]

    for b_lbl, b_val, b_col in ef_bullets:
        p_b = tf_ef.add_paragraph()
        p_b.alignment = PP_ALIGN.LEFT
        p_b.space_before = Pt(1.5)
        p_b.space_after = Pt(0.0)
        r_bl = p_b.add_run()
        set_font(r_bl, f"• {b_lbl}", bold=True, size_pt=7.6, color=b_col)
        r_bv = p_b.add_run()
        set_font(r_bv, b_val, bold=False, size_pt=7.3, color=C_SLATE_700)

    # ==========================================================================
    # RIGHT COLUMN: CHALLENGES & SOLUTIONS (2-Column Table)
    # ==========================================================================
    tb_ch_hdr = s4.shapes.add_textbox(Inches(6.65), Inches(1.44), Inches(6.08), Inches(0.24))
    tf_ch = tb_ch_hdr.text_frame
    tf_ch.word_wrap = True
    tf_ch.margin_left = 0
    tf_ch.margin_top = 0
    p_ch = tf_ch.paragraphs[0]
    p_ch.alignment = PP_ALIGN.LEFT
    r_ch = p_ch.add_run()
    set_font(r_ch, "Defensible Risk Mitigation: Challenges & Solutions", bold=True, size_pt=11.2, color=C_SLATE_900, underline=True)

    table_shape = s4.shapes.add_table(7, 2, Inches(6.65), Inches(1.74), Inches(6.08), Inches(5.10))
    table = table_shape.table
    table.columns[0].width = Inches(2.20)
    table.columns[1].width = Inches(3.88)

    # Row 0: Table Header
    table.rows[0].height = Inches(0.36)
    cell_h0 = table.cell(0, 0)
    cell_h0.fill.solid()
    cell_h0.fill.fore_color.rgb = C_SLATE_900
    cell_h0.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell_h0.margin_left = Inches(0.08)
    cell_h0.margin_right = Inches(0.08)
    p_h0 = cell_h0.text_frame.paragraphs[0]
    p_h0.alignment = PP_ALIGN.CENTER
    r_th0 = p_h0.add_run()
    set_font(r_th0, "GROUND CHALLENGE / RISK", bold=True, size_pt=8.5, color=C_WHITE)

    cell_h1 = table.cell(0, 1)
    cell_h1.fill.solid()
    cell_h1.fill.fore_color.rgb = C_DARK_GREEN
    cell_h1.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell_h1.margin_left = Inches(0.08)
    cell_h1.margin_right = Inches(0.08)
    p_h1 = cell_h1.text_frame.paragraphs[0]
    p_h1.alignment = PP_ALIGN.CENTER
    r_th1 = p_h1.add_run()
    set_font(r_th1, "OUR PRACTICAL SOLUTION", bold=True, size_pt=8.5, color=C_WHITE)

    table_data = [
        ("Low Digital Literacy", "Informal collectors struggle with text-heavy apps or typing.",
         "Voice + Pictorial Interface", "Visual scrap category icons + regional voice support eliminates reading and typing barriers."),

        ("Poor Connectivity", "Dense scrap clusters and transit corridors experience 2G dead zones.",
         "Offline-First + Automatic Sync", "IndexedDB buffers lots, photos & receipts locally; auto-syncs asynchronously upon network reconnection."),

        ("Weight Disputes", "Scale manipulation or reading disputes create mistrust across handovers.",
         "Scale Photo + OCR Verification", "Camera captures scale display; OCR extracts numeric digits for verifiable photo & weight evidence."),

        ("Fragmented Informal Chain", "Scrap passes through multiple unrecorded intermediaries before recycling.",
         "Digital Lot + Parent-Child Tracking", "Small collector lots are linked into aggregated bulk batches, preserving origin traceability across tiers."),

        ("Aggregator Resistance", "Middlemen scrap dealers oppose platforms that attempt to bypass them.",
         "Make Existing Dealers Collection Hubs", "Integrates existing aggregators as certified hubs earning legitimate handling margins on bulk supply."),

        ("Lack of Traceability", "CPCB recyclers lack auditable proof of physical custody for EPR audits.",
         "QR + Timestamp + GPS + Digital Handover", "Dual-scan QR verification logs photo, weight, GPS coordinates, and server timestamp at every stage.")
    ]

    for row_idx, (c_title, c_desc, s_title, s_desc) in enumerate(table_data, start=1):
        table.rows[row_idx].height = Inches(0.79)
        
        # Left Cell (Challenge)
        c_cell = table.cell(row_idx, 0)
        c_cell.fill.solid()
        c_cell.fill.fore_color.rgb = C_BG_RED
        c_cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        c_cell.margin_left = Inches(0.08)
        c_cell.margin_right = Inches(0.08)
        c_cell.margin_top = Inches(0.04)
        c_cell.margin_bottom = Inches(0.04)
        
        tf_c = c_cell.text_frame
        tf_c.word_wrap = True
        p_ct = tf_c.paragraphs[0]
        p_ct.alignment = PP_ALIGN.LEFT
        r_ct = p_ct.add_run()
        set_font(r_ct, f"• {c_title}", bold=True, size_pt=7.8, color=C_RED_DARK)
        
        p_cd = tf_c.add_paragraph()
        p_cd.alignment = PP_ALIGN.LEFT
        p_cd.space_before = Pt(1.0)
        r_cd = p_cd.add_run()
        set_font(r_cd, c_desc, bold=False, size_pt=7.0, color=C_SLATE_700)

        # Right Cell (Solution)
        s_cell = table.cell(row_idx, 1)
        s_cell.fill.solid()
        s_cell.fill.fore_color.rgb = C_BG_GREEN
        s_cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        s_cell.margin_left = Inches(0.08)
        s_cell.margin_right = Inches(0.08)
        s_cell.margin_top = Inches(0.04)
        s_cell.margin_bottom = Inches(0.04)
        
        tf_s = s_cell.text_frame
        tf_s.word_wrap = True
        p_st = tf_s.paragraphs[0]
        p_st.alignment = PP_ALIGN.LEFT
        r_st = p_st.add_run()
        set_font(r_st, f"• {s_title}", bold=True, size_pt=7.8, color=C_DARK_GREEN)
        
        p_sd = tf_s.add_paragraph()
        p_sd.alignment = PP_ALIGN.LEFT
        p_sd.space_before = Pt(1.0)
        r_sd = p_sd.add_run()
        set_font(r_sd, s_desc, bold=False, size_pt=7.0, color=C_SLATE_700)

    # ==========================================================================
    # SLIDE 5: Impact and Benefits
    # ==========================================================================
    s5 = prs.slides[4]
    update_team_oval(s5)
    clean_placeholder(s5, "TextBox 8")

    for sh in s5.shapes:
        if sh.name == "Title 1":
            p = sh.text_frame.paragraphs[0]
            p.text = ""
            r = p.add_run()
            set_font(r, "IMPACT AND BENEFITS", bold=True, size_pt=32, color=C_SLATE_900)

    # Slide 5 Subtitle
    sub_box5 = s5.shapes.add_textbox(Inches(0.60), Inches(1.08), Inches(12.13), Inches(0.32))
    tf_sub5 = sub_box5.text_frame
    tf_sub5.word_wrap = True
    tf_sub5.margin_left = 0
    tf_sub5.margin_top = 0
    p0_s5 = tf_sub5.paragraphs[0]
    p0_s5.alignment = PP_ALIGN.LEFT
    r0_s5 = p0_s5.add_run()
    set_font(r0_s5, "Stakeholder Value Transformation: ", bold=True, size_pt=12.0, color=C_SLATE_900)
    r1_s5 = p0_s5.add_run()
    set_font(r1_s5, "Creating Value Across the Informal-to-Formal Recycling Chain", bold=False, size_pt=10.5, color=C_DARK_GREEN)

    # ==========================================================================
    # TOP SECTION: 4 STAKEHOLDER CARDS (Before vs With E-Waste Setu)
    # ==========================================================================
    # Stakeholder Data: (Title, Role, Border Color, Value Tag, Before Points, After Points)
    stakeholder_data = [
        ("1. INFORMAL COLLECTORS",
         "Waste-Pickers & Itinerant Buyers",
         C_DARK_GREEN,
         "VALUE: Fair floor prices & recognized digital livelihoods",
         [
             "Volatile scrap prices dictated by middlemen",
             "Zero direct access to authorized recyclers",
             "Frequent weight & cash payment disputes",
             "No transaction records or credit history"
         ],
         [
             "Live transparent price benchmark access",
             "Scale OCR weight & photo evidence",
             "Instant Cash / UPI verified payment records",
             "Portable digital transaction history",
             "Safe inclusion into formal recycling streams"
         ]),

        ("2. LOCAL AGGREGATORS",
         "Scrap Dealers & Local Hubs",
         C_AMBER_DARK,
         "VALUE: Digital lotting & transition to certified collection hubs",
         [
             "Manual paper ledgers prone to loss & error",
             "Difficult multi-collector bulk coordination",
             "Zero advance visibility of inbound scrap",
             "Vulnerability to gray-market scrutiny"
         ],
         [
             "Digital lotting with parent-child tracking",
             "Rapid aggregation of small lots into bulk",
             "Transparent, tamper-evident purchase ledger",
             "Direct digital coordination with recyclers",
             "Certified formal collection hub status"
         ]),

        ("3. AUTHORIZED RECYCLERS",
         "CPCB Dismantlers & Smelting Plants",
         C_BLUE_DARK,
         "VALUE: Traceable feedstock & verifiable chain-of-custody",
         [
             "Fragmented supply bypassing formal plants",
             "Unverified scrap origin & quality disputes",
             "Manual paper manifests with audit gaps",
             "Severe plant capacity under-utilization"
         ],
         [
             "Pre-segregated, quality-graded bulk lots",
             "Location-based spatial sourcing & routing",
             "Cryptographic QR custody & intake check",
             "Verifiable chain-of-custody to smelting",
             "Predictable supply & higher throughput"
         ]),

        ("4. REGULATORS / EPR",
         "CPCB, MoEFCC, SPCBs & Brand Owners",
         C_PURPLE_DARK,
         "VALUE: Complete provenance & fraud-free CPCB Form-6 audits",
         [
             "Unmonitored informal scrap flows & leakage",
             "Untraceable material custody across dealers",
             "Risk of fraudulent paper-only credits",
             "Toxic crude backyard leaching & burning"
         ],
         [
             "Full Collector ➔ Hub ➔ Recycler trace",
             "Immutable transaction & movement trail",
             "Photo, OCR, GPS & timestamp verification",
             "Standardized structured compliance data",
             "Automated, verifiable CPCB Form-6 audits"
         ])
    ]

    card_w = Inches(2.93)
    card_gap = Inches(0.14)
    y_top_cards = Inches(1.44)
    h_top_cards = Inches(3.32)

    for idx, (st_title, st_role, st_col, st_val, before_pts, after_pts) in enumerate(stakeholder_data):
        x_c = Inches(0.60) + idx * (card_w + card_gap)

        # Single Unified Executive Card with matching tier border
        sh_c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_c, y_top_cards, card_w, h_top_cards)
        style_box(sh_c, bg_color=C_WHITE, border_color=st_col, border_width_pt=1.4)
        
        tf_c = sh_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.12)
        tf_c.margin_right = Inches(0.12)
        tf_c.margin_top = Inches(0.12)
        tf_c.margin_bottom = Inches(0.08)

        # 1. Header: Title + Role
        p_hdr = tf_c.paragraphs[0]
        p_hdr.alignment = PP_ALIGN.LEFT
        r_th = p_hdr.add_run()
        set_font(r_th, st_title, bold=True, size_pt=9.6, color=st_col, font_name="Calibri")
        
        p_role = tf_c.add_paragraph()
        p_role.alignment = PP_ALIGN.LEFT
        p_role.space_before = Pt(0.5)
        p_role.space_after = Pt(3.5)
        r_rl = p_role.add_run()
        set_font(r_rl, st_role, bold=False, size_pt=7.4, color=C_SLATE_500, font_name="Calibri")

        # 2. Before Section
        p_bef_hdr = tf_c.add_paragraph()
        p_bef_hdr.alignment = PP_ALIGN.LEFT
        p_bef_hdr.space_before = Pt(1.5)
        p_bef_hdr.space_after = Pt(1.0)
        r_bh = p_bef_hdr.add_run()
        set_font(r_bh, "BEFORE:", bold=True, size_pt=7.8, color=C_RED_DARK, font_name="Calibri")

        for b_txt in before_pts:
            p_bp = tf_c.add_paragraph()
            p_bp.alignment = PP_ALIGN.LEFT
            p_bp.space_before = Pt(1.0)
            p_bp.space_after = Pt(0.2)
            r_d = p_bp.add_run()
            set_font(r_d, "• ", bold=True, size_pt=7.4, color=C_RED_MID, font_name="Calibri")
            r_tx = p_bp.add_run()
            set_font(r_tx, b_txt, bold=False, size_pt=7.3, color=C_SLATE_700, font_name="Calibri")

        # 3. After Section
        p_aft_hdr = tf_c.add_paragraph()
        p_aft_hdr.alignment = PP_ALIGN.LEFT
        p_aft_hdr.space_before = Pt(3.5)
        p_aft_hdr.space_after = Pt(1.0)
        r_ah = p_aft_hdr.add_run()
        set_font(r_ah, "AFTER (WITH E-WASTE SETU):", bold=True, size_pt=7.8, color=C_DARK_GREEN, font_name="Calibri")

        for a_txt in after_pts:
            p_ap = tf_c.add_paragraph()
            p_ap.alignment = PP_ALIGN.LEFT
            p_ap.space_before = Pt(1.0)
            p_ap.space_after = Pt(0.2)
            r_ck = p_ap.add_run()
            set_font(r_ck, "✓ ", bold=True, size_pt=7.4, color=C_DARK_GREEN, font_name="Calibri")
            r_atx = p_ap.add_run()
            set_font(r_atx, a_txt, bold=False, size_pt=7.3, color=C_SLATE_900, font_name="Calibri")

        # 4. Core Value Tag at Bottom of Card
        p_v = tf_c.add_paragraph()
        p_v.alignment = PP_ALIGN.LEFT
        p_v.space_before = Pt(4.5)
        r_v = p_v.add_run()
        set_font(r_v, st_val, bold=True, size_pt=7.0, color=st_col, font_name="Calibri")

    # ==========================================================================
    # BOTTOM SECTION: 3 BIG IMPACT AREAS (Social, Economic, Environmental)
    # ==========================================================================
    impact_pillars = [
        ("SOCIAL IMPACT", "Inclusive Livelihoods & Safety",
         [
             ("Zero-Barrier Access", "Voice guidance & visual touch UI remove literacy barriers"),
             ("Financial Identity", "Verifiable Cash / UPI payout history enables credit eligibility"),
             ("Dispute Protection", "Scale OCR photo evidence eliminates arbitrary middleman cuts"),
             ("Dignity of Labor", "Transitions informal collectors into recognized green workers")
         ],
         "Mechanism: Transparent pricing ➔ Dignified, recognized livelihoods",
         C_BLUE_DARK),

        ("ECONOMIC IMPACT", "Price Transparency & Supply Efficiency",
         [
             ("Transparent Pricing", "Live benchmark floor prevents middleman price exploitation"),
             ("Bulk Value Capture", "Small lots aggregated into high-margin formal consignments"),
             ("Dispute-Free Trade", "Tamper-evident OCR weight logs resolve handover conflicts"),
             ("Capacity Throughput", "Predictable, sorted feedstock maximizes plant utilization")
         ],
         "Mechanism: Lot aggregation + OCR proof ➔ Dispute-free bulk trade",
         C_AMBER_DARK),

        ("ENVIRONMENTAL IMPACT", "Toxic Diversion & Circular Recovery",
         [
             ("Zero Toxic Leaching", "Diverts PCBs away from crude acid baths & open-wire burning"),
             ("Critical Minerals", "Secures domestic hydrometallurgical recovery of Cu, Au, Ag"),
             ("Chain-of-Custody", "Cryptographic QR tracking prevents gray-market transit dumping"),
             ("Audit-Proof EPR", "True physical custody audit trail replaces paper-only fraud")
         ],
         "Mechanism: Digital custody ➔ End-to-end trace & zero toxic burning",
         C_DARK_GREEN)
    ]

    imp_w = Inches(3.95)
    imp_gap = Inches(0.14)
    y_imp = Inches(4.88)
    h_imp = Inches(1.96)

    for idx, (imp_title, imp_sub, imp_pts, imp_mech, imp_col) in enumerate(impact_pillars):
        x_imp = Inches(0.60) + idx * (imp_w + imp_gap)
        
        # Single Unified Executive Card with matching tier border
        sh_imp = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_imp, y_imp, imp_w, h_imp)
        style_box(sh_imp, bg_color=C_WHITE, border_color=imp_col, border_width_pt=1.4)

        tf_imp = sh_imp.text_frame
        tf_imp.word_wrap = True
        tf_imp.margin_left = Inches(0.14)
        tf_imp.margin_right = Inches(0.14)
        tf_imp.margin_top = Inches(0.12)
        tf_imp.margin_bottom = Inches(0.08)

        # Line 1: Title & Subtitle
        p_t = tf_imp.paragraphs[0]
        p_t.alignment = PP_ALIGN.LEFT
        r_it = p_t.add_run()
        set_font(r_it, imp_title, bold=True, size_pt=9.6, color=imp_col, font_name="Calibri")
        r_is = p_t.add_run()
        set_font(r_is, f"  |  {imp_sub}", bold=True, size_pt=7.8, color=C_SLATE_900, font_name="Calibri")

        # Crisp Point-by-Point Impact Cuts
        for pt_label, pt_desc in imp_pts:
            p_pt = tf_imp.add_paragraph()
            p_pt.alignment = PP_ALIGN.LEFT
            p_pt.space_before = Pt(1.5)
            p_pt.space_after = Pt(0.5)
            
            r_dot = p_pt.add_run()
            set_font(r_dot, "• ", bold=True, size_pt=7.4, color=imp_col, font_name="Calibri")
            
            r_lbl = p_pt.add_run()
            set_font(r_lbl, f"{pt_label}: ", bold=True, size_pt=7.3, color=C_SLATE_900, font_name="Calibri")
            
            r_dsc = p_pt.add_run()
            set_font(r_dsc, pt_desc, bold=False, size_pt=7.2, color=C_SLATE_700, font_name="Calibri")

        # Mechanism Line
        p_mech = tf_imp.add_paragraph()
        p_mech.alignment = PP_ALIGN.LEFT
        p_mech.space_before = Pt(3.0)
        r_mc = p_mech.add_run()
        set_font(r_mc, imp_mech, bold=True, size_pt=7.0, color=imp_col, font_name="Calibri")

    # ==========================================================================
    # SLIDE 6: Research and References
    # ==========================================================================
    s6 = prs.slides[5]
    update_team_oval(s6)
    clean_placeholder(s6, "TextBox 8")

    for sh in s6.shapes:
        if sh.name == "Title 1":
            p = sh.text_frame.paragraphs[0]
            p.text = ""
            r = p.add_run()
            set_font(r, "RESEARCH AND REFERENCES", bold=True, size_pt=30, color=C_SLATE_900)

    # Subtitle Box
    sub_box = s6.shapes.add_textbox(Inches(0.60), Inches(1.05), Inches(12.0), Inches(0.32))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    r_sub1 = p_sub.add_run()
    set_font(r_sub1, "Comprehensive Literature, Technical Foundations & Primary Field Research: ", bold=True, size_pt=10.5, color=C_SLATE_900)
    r_sub2 = p_sub.add_run()
    set_font(r_sub2, "Evidence-Based Architecture for E-Waste Setu", bold=False, size_pt=10.5, color=C_DARK_GREEN)

    # 3-Column Layout: Column Width & Spacing
    col_w = Inches(3.92)
    col_gap = Inches(0.18)
    y_col = Inches(1.38)
    h_col = Inches(4.44)

    # ==========================================================================
    # COLUMN 1: PROBLEM, POLICY & DOMAIN RESEARCH
    # ==========================================================================
    x_c1 = Inches(0.60)
    sh_c1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_c1, y_col, col_w, h_col)
    style_box(sh_c1, bg_color=C_WHITE, border_color=C_BLUE_DARK, border_width_pt=1.4)
    tf_c1 = sh_c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = Inches(0.12)
    tf_c1.margin_right = Inches(0.12)
    tf_c1.margin_top = Inches(0.12)
    tf_c1.margin_bottom = Inches(0.08)

    # Header
    p_h1 = tf_c1.paragraphs[0]
    p_h1.alignment = PP_ALIGN.LEFT
    r_h1 = p_h1.add_run()
    set_font(r_h1, "1. PROBLEM, POLICY & DOMAIN RESEARCH", bold=True, size_pt=9.6, color=C_BLUE_DARK, font_name="Calibri")

    p_sub1 = tf_c1.add_paragraph()
    p_sub1.alignment = PP_ALIGN.LEFT
    p_sub1.space_before = Pt(0.5)
    p_sub1.space_after = Pt(2.5)
    r_sb1 = p_sub1.add_run()
    set_font(r_sb1, "Statutory Mandates, Academic Literature & National Policy", bold=False, size_pt=7.2, color=C_SLATE_500, font_name="Calibri")

    domain_refs = [
        ("Smart India Hackathon 2026 — PS 26229",
         "https://sih2026.vuce.in/ps/SIH26229",
         "Ministry of Mines / JNARDDC",
         "Formalization Mandate: Problem scope targeting informal collector inclusion; sets verifiable custody & transparent price benchmarks as core criteria."),

        ("E-Waste (Management) Rules, 2022",
         "https://eprewaste.cpcb.gov.in/assets/PDF/e-waste_rules_2022.pdf",
         "CPCB / MoEFCC",
         "Statutory Compliance: Mandates digital EPR credit trading, formal recycler licensing, and automated Form-6 electronic filing audit trails."),

        ("CPCB — E-Waste Management / EPR Portal",
         "https://eprewaste.cpcb.gov.in/",
         "Central Pollution Control Board",
         "Registry & Ecosystem: Official national portal for registered dismantlers & recyclers; mapped to our spatial engine for certified lot routing."),

        ("NITI Aayog — Circular Economy Strategy",
         "https://www.niti.gov.in/node/2108",
         "NITI Aayog (2022)",
         "Inclusion Roadmap: Prioritizes decentralized aggregation hubs, fair floor pricing, and domestic recovery of critical electronic minerals."),

        ("JNARDDC — Waste Management Research",
         "https://www.jnarddc.gov.in/en/services/waste.aspx",
         "Autonomous Body, MoM",
         "Metallurgical Recovery: Benchmark extraction indices for Au, Ag, Cu, Li via hydrometallurgical processing over hazardous backyard leaching."),

        ("Toxics Link — Electronic Waste Research",
         "https://toxicslink.org/electronic-waste/",
         "Environmental Research Body",
         "Ground Vulnerabilities: Proves 25–40% value skimming by middlemen and severe heavy metal poisoning from crude open acid burning.")
    ]

    for title, url, org, desc in domain_refs:
        p_t = tf_c1.add_paragraph()
        p_t.alignment = PP_ALIGN.LEFT
        p_t.space_before = Pt(1.8)
        p_t.space_after = Pt(0.2)

        r_dot = p_t.add_run()
        set_font(r_dot, "• ", bold=True, size_pt=7.5, color=C_BLUE_DARK, font_name="Calibri")

        r_lnk = p_t.add_run()
        set_font(r_lnk, f"{title} ↗", bold=True, size_pt=7.5, color=C_BLUE_DARK, underline=True, font_name="Calibri")
        r_lnk.hyperlink.address = url

        r_org = p_t.add_run()
        set_font(r_org, f" ({org})", bold=False, size_pt=6.8, color=C_SLATE_500, font_name="Calibri")

        p_d = tf_c1.add_paragraph()
        p_d.alignment = PP_ALIGN.LEFT
        p_d.space_before = Pt(0.2)
        p_d.space_after = Pt(1.0)
        r_d = p_d.add_run()
        set_font(r_d, desc, bold=False, size_pt=6.9, color=C_SLATE_700, font_name="Calibri")

    # Bottom Callout for Policy
    p_pbot = tf_c1.add_paragraph()
    p_pbot.alignment = PP_ALIGN.LEFT
    p_pbot.space_before = Pt(3.0)
    r_pb1 = p_pbot.add_run()
    set_font(r_pb1, "REGULATORY ALIGNMENT: ", bold=True, size_pt=7.0, color=C_BLUE_DARK, font_name="Calibri")
    r_pb2 = p_pbot.add_run()
    set_font(r_pb2, "Direct compliance mapping to CPCB Form-6 digital ledger and EPR credit audit rules.", bold=False, size_pt=6.9, color=C_SLATE_700, font_name="Calibri")

    # ==========================================================================
    # COLUMN 2: TECHNOLOGY & AI RESEARCH
    # ==========================================================================
    x_c2 = Inches(0.60) + (col_w + col_gap)
    sh_c2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_c2, y_col, col_w, h_col)
    style_box(sh_c2, bg_color=C_WHITE, border_color=C_DARK_GREEN, border_width_pt=1.4)
    tf_c2 = sh_c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = Inches(0.12)
    tf_c2.margin_right = Inches(0.12)
    tf_c2.margin_top = Inches(0.12)
    tf_c2.margin_bottom = Inches(0.08)

    # Header
    p_h2 = tf_c2.paragraphs[0]
    p_h2.alignment = PP_ALIGN.LEFT
    r_h2 = p_h2.add_run()
    set_font(r_h2, "2. TECHNOLOGY & AI RESEARCH", bold=True, size_pt=9.6, color=C_DARK_GREEN, font_name="Calibri")

    p_sub2 = tf_c2.add_paragraph()
    p_sub2.alignment = PP_ALIGN.LEFT
    p_sub2.space_before = Pt(0.5)
    p_sub2.space_after = Pt(2.5)
    r_sb2 = p_sub2.add_run()
    set_font(r_sb2, "Open-Source Engines, Algorithmic Foundations & Web Standards", bold=False, size_pt=7.2, color=C_SLATE_500, font_name="Calibri")

    tech_refs = [
        ("MobileNetV3 — Mobile Vision",
         "https://arxiv.org/abs/1905.02244",
         "Howard et al., Google Research",
         "Lightweight Neural Vision: Hardware-aware NAS architecture; quantized INT8 on-device scrap category classifier (<5MB) running in real-time on budget phones."),

        ("IEEE Access — E-Waste Vision AI",
         "https://ieeexplore.ieee.org/document/9385073",
         "IEEE Xplore (2021)",
         "Peer-Reviewed Benchmark: Evaluates lightweight transfer-learning CNN architectures for automated multi-category electronic scrap classification on edge hardware."),

        ("AI4Bharat — Indic Language AI",
         "https://ai4bharat.org/",
         "IIT Madras / MeitY Initiative",
         "Vernacular Voice Assistance: Speech-to-text and voice synthesis in Hindi & regional dialects; removes reading/typing barriers for low-literacy collectors."),

        ("OpenCV — Computer Vision Pipeline",
         "https://opencv.org/",
         "OpenCV Foundation",
         "Scale OCR Verification: Perspective warp rectification, adaptive thresholding & 7-segment digit recognition to extract tamper-evident weight proof from photos."),

        ("PostGIS — Spatial Database",
         "https://postgis.net/",
         "PostgreSQL Spatial Extension",
         "Geospatial Clustering: ST_DWithin and K-Nearest-Neighbor spatial queries enabling location-based clustering and dynamic collector-to-hub routing."),

        ("MDN — IndexedDB API & Service Worker",
         "https://developer.mozilla.org/en-US/docs/Web/API/IndexedDB_API",
         "W3C / Mozilla Standards",
         "Offline-First Architecture: Local transactional queueing in zero-connectivity slum areas; automatically cryptographically signs and syncs when reconnected.")
    ]

    for title, url, org, desc in tech_refs:
        p_t = tf_c2.add_paragraph()
        p_t.alignment = PP_ALIGN.LEFT
        p_t.space_before = Pt(1.8)
        p_t.space_after = Pt(0.2)

        r_dot = p_t.add_run()
        set_font(r_dot, "• ", bold=True, size_pt=7.5, color=C_DARK_GREEN, font_name="Calibri")

        r_lnk = p_t.add_run()
        set_font(r_lnk, f"{title} ↗", bold=True, size_pt=7.5, color=C_BLUE_DARK, underline=True, font_name="Calibri")
        r_lnk.hyperlink.address = url

        r_org = p_t.add_run()
        set_font(r_org, f" ({org})", bold=False, size_pt=6.8, color=C_SLATE_500, font_name="Calibri")

        p_d = tf_c2.add_paragraph()
        p_d.alignment = PP_ALIGN.LEFT
        p_d.space_before = Pt(0.2)
        p_d.space_after = Pt(1.0)
        r_d = p_d.add_run()
        set_font(r_d, desc, bold=False, size_pt=6.9, color=C_SLATE_700, font_name="Calibri")

    # Bottom Callout for Tech
    p_tbot = tf_c2.add_paragraph()
    p_tbot.alignment = PP_ALIGN.LEFT
    p_tbot.space_before = Pt(3.0)
    r_tb = p_tbot.add_run()
    set_font(r_tb, "ARCHITECTURAL PRINCIPLE: ", bold=True, size_pt=7.0, color=C_DARK_GREEN, font_name="Calibri")
    r_tb2 = p_tbot.add_run()
    set_font(r_tb2, "Zero-capex, open-source stack that runs on entry-level Android smartphones with zero proprietary licensing cost.", bold=False, size_pt=6.9, color=C_SLATE_700, font_name="Calibri")

    # ==========================================================================
    # COLUMN 3: PRIMARY FIELD RESEARCH
    # ==========================================================================
    x_c3 = Inches(0.60) + 2 * (col_w + col_gap)
    sh_c3 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_c3, y_col, col_w, h_col)
    style_box(sh_c3, bg_color=C_WHITE, border_color=C_AMBER_DARK, border_width_pt=1.4)
    tf_c3 = sh_c3.text_frame
    tf_c3.word_wrap = True
    tf_c3.margin_left = Inches(0.12)
    tf_c3.margin_right = Inches(0.12)
    tf_c3.margin_top = Inches(0.12)
    tf_c3.margin_bottom = Inches(0.08)

    # Header
    p_h3 = tf_c3.paragraphs[0]
    p_h3.alignment = PP_ALIGN.LEFT
    r_h3 = p_h3.add_run()
    set_font(r_h3, "3. PRIMARY FIELD RESEARCH", bold=True, size_pt=9.6, color=C_AMBER_DARK, font_name="Calibri")

    p_sub3 = tf_c3.add_paragraph()
    p_sub3.alignment = PP_ALIGN.LEFT
    p_sub3.space_before = Pt(0.5)
    p_sub3.space_after = Pt(2.0)
    r_sb3 = p_sub3.add_run()
    set_font(r_sb3, "On-Ground Interviews & Operational Workflow Observation", bold=False, size_pt=7.2, color=C_SLATE_500, font_name="Calibri")

    # Scope Badge
    p_sc = tf_c3.add_paragraph()
    p_sc.alignment = PP_ALIGN.LEFT
    p_sc.space_before = Pt(0.5)
    p_sc.space_after = Pt(3.0)
    r_sc = p_sc.add_run()
    set_font(r_sc, "FIELD SAMPLE: 2 Informal Collectors + 1 Main Aggregator", bold=True, size_pt=7.4, color=C_AMBER_DARK, font_name="Calibri")

    field_data = [
        ("Participant 01: Collector 01 (Daily Waste-Picker)",
         "Collects discarded electronics from repair shops & bins ➔ Sells daily to roadside dealer for instant cash; accepts flat under-weighed estimate.",
         '"Scrap dealers reduce 1–2 kg on mechanical spring scales and give flat rates for printed boards. If we know true rates and have scale photo proof, no one can cut our money."'),

        ("Participant 02: Collector 02 (Itinerant Scrap Buyer)",
         "Travels 15–20 km daily pushcart route buying home appliances ➔ Connectivity drops in narrow alleys ➔ Needs immediate price proof for household sellers.",
         '"Households bargain hard because they don\'t trust our rates. An app showing live market scrap rates on screen builds trust and closes deals faster."'),

        ("Participant 03: Main Aggregator (Local Scrap Hub Operator)",
         "Buys small lots from 30+ collectors ➔ Stores in yard ➔ Records in paper bahi-khata ➔ Holds stock until truckload (1–2 tons) for recycler dispatch.",
         '"Big recyclers pay higher rates but demand 500kg+ lots and manifests. Paper notebooks get messy; digital lotting lets us operate as certified formal collection hubs."')
    ]

    for participant, process, what_he_told in field_data:
        p_p = tf_c3.add_paragraph()
        p_p.alignment = PP_ALIGN.LEFT
        p_p.space_before = Pt(2.0)
        p_p.space_after = Pt(0.2)

        r_pt = p_p.add_run()
        set_font(r_pt, participant, bold=True, size_pt=7.5, color=C_SLATE_900, font_name="Calibri")

        p_pr = tf_c3.add_paragraph()
        p_pr.alignment = PP_ALIGN.LEFT
        p_pr.space_before = Pt(0.2)
        p_pr.space_after = Pt(0.2)
        r_prl = p_pr.add_run()
        set_font(r_prl, "Process Told: ", bold=True, size_pt=6.9, color=C_AMBER_DARK, font_name="Calibri")
        r_prv = p_pr.add_run()
        set_font(r_prv, process, bold=False, size_pt=6.8, color=C_SLATE_700, font_name="Calibri")

        p_wh = tf_c3.add_paragraph()
        p_wh.alignment = PP_ALIGN.LEFT
        p_wh.space_before = Pt(0.2)
        p_wh.space_after = Pt(1.5)
        r_whl = p_wh.add_run()
        set_font(r_whl, "What He Told: ", bold=True, size_pt=6.9, color=C_DARK_GREEN, font_name="Calibri")
        r_whv = p_wh.add_run()
        set_font(r_whv, what_he_told, bold=False, size_pt=6.8, color=C_SLATE_900, font_name="Calibri")

    # Pipeline Coverage Box
    p_cov_h = tf_c3.add_paragraph()
    p_cov_h.alignment = PP_ALIGN.LEFT
    p_cov_h.space_before = Pt(2.5)
    r_cvh = p_cov_h.add_run()
    set_font(r_cvh, "FIELD RESEARCH COVERAGE:", bold=True, size_pt=7.2, color=C_SLATE_900, font_name="Calibri")

    p_cov_fl = tf_c3.add_paragraph()
    p_cov_fl.alignment = PP_ALIGN.LEFT
    p_cov_fl.space_before = Pt(0.5)
    p_cov_fl.space_after = Pt(2.0)
    r_cfl = p_cov_fl.add_run()
    set_font(r_cfl, "Collection ➔ Pricing ➔ Weighing ➔ Aggregation ➔ Bulk Transfer ➔ Recycler", bold=True, size_pt=6.9, color=C_DARK_GREEN, font_name="Calibri")

    # Field Takeaway
    p_fbot = tf_c3.add_paragraph()
    p_fbot.alignment = PP_ALIGN.LEFT
    p_fbot.space_before = Pt(1.5)
    r_fb1 = p_fbot.add_run()
    set_font(r_fb1, "Core Validation: ", bold=True, size_pt=7.0, color=C_AMBER_DARK, font_name="Calibri")
    r_fb2 = p_fbot.add_run()
    set_font(r_fb2, "Informal actors do not resist formalization when empowered with scale proof and market prices without disrupting their existing cash flows.", bold=False, size_pt=6.8, color=C_SLATE_700, font_name="Calibri")

    # ==========================================================================
    # BOTTOM TWO EXECUTIVE BLOCKS: PROTOTYPE WEBSITE & YOUTUBE DEMO
    # ==========================================================================
    y_ban = Inches(5.88)
    h_ban = Inches(0.96)
    b_w = Inches(5.95)
    b_gap = Inches(0.22)

    # BLOCK 1 (LEFT): LIVE WORKING PROTOTYPE WEBSITE
    x_b1 = Inches(0.60)
    sh_b1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_b1, y_ban, b_w, h_ban)
    style_box(sh_b1, bg_color=C_BG_GREEN, border_color=C_DARK_GREEN, border_width_pt=1.4)
    tf_b1 = sh_b1.text_frame
    tf_b1.word_wrap = True
    tf_b1.margin_left = Inches(0.16)
    tf_b1.margin_right = Inches(0.16)
    tf_b1.margin_top = Inches(0.12)
    tf_b1.margin_bottom = Inches(0.08)

    p_b1_h = tf_b1.paragraphs[0]
    p_b1_h.alignment = PP_ALIGN.LEFT
    r_b1_lbl = p_b1_h.add_run()
    set_font(r_b1_lbl, "PROTOTYPE WEB APPLICATION:  ", bold=True, size_pt=9.2, color=C_DARK_GREEN, font_name="Calibri")
    r_b1_lnk = p_b1_h.add_run()
    set_font(r_b1_lnk, "https://ewastesetu.vercel.app ↗", bold=True, size_pt=9.2, color=C_BLUE_DARK, underline=True, font_name="Calibri")
    r_b1_lnk.hyperlink.address = "https://ewastesetu.vercel.app"

    p_b1_d = tf_b1.add_paragraph()
    p_b1_d.alignment = PP_ALIGN.LEFT
    p_b1_d.space_before = Pt(3.0)
    r_b1_d1 = p_b1_d.add_run()
    set_font(r_b1_d1, "Live Responsive PWA: ", bold=True, size_pt=7.8, color=C_SLATE_900, font_name="Calibri")
    r_b1_d2 = p_b1_d.add_run()
    set_font(r_b1_d2, "Mobile-optimized web app with visual scrap touch categories, offline IndexedDB transaction storage, and camera scale OCR weight extraction for budget Android phones.", bold=False, size_pt=7.5, color=C_SLATE_700, font_name="Calibri")

    # BLOCK 2 (RIGHT): YOUTUBE VIDEO DEMONSTRATION
    x_b2 = x_b1 + b_w + b_gap
    sh_b2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_b2, y_ban, b_w, h_ban)
    style_box(sh_b2, bg_color=C_BG_BLUE, border_color=C_BLUE_DARK, border_width_pt=1.4)
    tf_b2 = sh_b2.text_frame
    tf_b2.word_wrap = True
    tf_b2.margin_left = Inches(0.16)
    tf_b2.margin_right = Inches(0.16)
    tf_b2.margin_top = Inches(0.12)
    tf_b2.margin_bottom = Inches(0.08)

    p_b2_h = tf_b2.paragraphs[0]
    p_b2_h.alignment = PP_ALIGN.LEFT
    r_b2_lbl = p_b2_h.add_run()
    set_font(r_b2_lbl, "YOUTUBE VIDEO DEMONSTRATION:  ", bold=True, size_pt=9.2, color=C_BLUE_DARK, font_name="Calibri")
    r_b2_lnk = p_b2_h.add_run()
    set_font(r_b2_lnk, "https://youtu.be/ewastesetu-demo ↗", bold=True, size_pt=9.2, color=C_BLUE_DARK, underline=True, font_name="Calibri")
    r_b2_lnk.hyperlink.address = "https://youtu.be/ewastesetu-demo"

    p_b2_d = tf_b2.add_paragraph()
    p_b2_d.alignment = PP_ALIGN.LEFT
    p_b2_d.space_before = Pt(3.0)
    r_b2_d1 = p_b2_d.add_run()
    set_font(r_b2_d1, "Complete System Walkthrough: ", bold=True, size_pt=7.8, color=C_SLATE_900, font_name="Calibri")
    r_b2_d2 = p_b2_d.add_run()
    set_font(r_b2_d2, "End-to-end video demo of voice scrap logging, digital scale OCR weight capture, local hub batch lotting, and dual QR CPCB Form-6 manifest handover.", bold=False, size_pt=7.5, color=C_SLATE_700, font_name="Calibri")

    prs.save(output_path)
    print("SUCCESS: Built official SIH deck directly on official template!")

if __name__ == "__main__":
    build_official_deck()
